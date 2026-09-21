"""The one-file workbook — the export for a device that cannot run a query.

Excel for iPad has no Power Query, so the CSV-plus-connection design has no
path on it at all. This file is the other shape, and these tests hold the
properties that make it usable on a tablet: every tab present even when
empty, headers frozen and filterable, blanks that stay blank, and numbers
that came from the export frames rather than from a second calculation.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from velocity.export.meta import ExportMeta
from velocity.export.workbook import (
    WORKBOOK_NAME,
    build_workbook,
    header_for,
)

TABS = ("Read Me", "Dashboard", "Bet Card", "Games", "Props",
        "Team Totals", "DFS Pool", "DFS Optimizer", "Curated Plays")
META = ExportMeta("2026-09-20T17:53:00Z", 2026, 3)


def _games() -> pd.DataFrame:
    return pd.DataFrame([{
        "game_id": "g1", "league": "nfl", "away_team": "Atlanta",
        "home_team": "Carolina", "market_spread": -2.5, "market_total": 41.0,
        "moneyline": -150, "model_away_score": 23.2, "model_home_score": 21.0,
        "model_total": 44.2, "spread_edge": -4.7, "total_edge": 3.2,
        "cover_probability": 0.421, "over_probability": 0.62,
        "confidence": 7.1, "kelly_fraction": 0.024,
        "weather": "51°F, wind 14 mph",
        "generated_at": META.generated_at, "season": 2026, "week": 3,
    }])


def _plays() -> pd.DataFrame:
    return pd.DataFrame([
        {"tier": "A+", "bet_type": "game", "selection": "Under 48.5",
         "market": "total", "edge": 0.10, "confidence": 7.4, "stake": 2.4,
         "reason": "Under 48.5 — Model total 44", "generated_at":
         META.generated_at, "season": 2026, "week": 3},
        {"tier": "Watch", "bet_type": "game", "selection": "Atlanta +2.5",
         "market": "spread", "edge": None, "confidence": 0.0, "stake": 0.0,
         "reason": "VETOED by the intel layer", "generated_at":
         META.generated_at, "season": 2026, "week": 3},
    ])


def _dfs() -> pd.DataFrame:
    return pd.DataFrame([{
        "player": "Bijan Robinson", "team": "ATL", "position": "RB",
        "salary": 7600, "projection": 17.4, "median": 16.2, "ceiling": 38.1,
        "ownership": None, "value_score": 2.29, "leverage_score": None,
        "stack_rating": 6.4, "generated_at": META.generated_at,
        "season": 2026, "week": 3,
    }])


def _dashboard() -> pd.DataFrame:
    return pd.DataFrame([
        {"section": "performance", "metric": "record", "value": None,
         "detail": "2-1-0"},
        {"section": "performance", "metric": "win_rate", "value": 0.6667,
         "detail": "decided bets only"},
        {"section": "roi", "metric": "roi_overall", "value": 0.2867,
         "detail": "profit per unit staked"},
        {"section": "clv", "metric": "mean_price_clv", "value": 0.0133,
         "detail": "the durable signal"},
        {"section": "top_plays", "metric": "play_1", "value": 0.10,
         "detail": "[A+] Under 48.5"},
    ])


def _build(tmp_path: Path, **frames: pd.DataFrame | None) -> Path:
    return build_workbook(tmp_path / WORKBOOK_NAME, META, **frames)


def test_every_tab_exists_even_with_nothing_to_put_in_it(tmp_path: Path) -> None:
    """An absent tab reads as a broken file; an empty one reads as an empty board."""
    wb = load_workbook(_build(tmp_path))
    assert wb.sheetnames == list(TABS)


def test_tab_order_matches_the_documented_workbook(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path, games=_games(), plays=_plays()))
    assert wb.sheetnames == list(TABS)
    assert wb.sheetnames[0] == "Read Me"


def test_headers_are_frozen_and_filterable(tmp_path: Path) -> None:
    """No mouse and no second window: tap to sort, and keep the header on screen."""
    wb = load_workbook(_build(tmp_path, games=_games()))
    ws = wb["Games"]
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref is not None
    assert ws.auto_filter.ref.startswith("A4:")


def test_the_metadata_columns_move_to_the_subtitle(tmp_path: Path) -> None:
    """Three constants repeated down every row is not information, it is width."""
    wb = load_workbook(_build(tmp_path, games=_games()))
    ws = wb["Games"]
    headers = [c.value for c in ws[4]]
    for gone in ("Generated At", "Season", "Week"):
        assert gone not in headers
    assert "2026-09-20T17:53:00Z" in str(ws["A2"].value)
    assert "Week 3" in str(ws["A2"].value)


def test_headers_are_readable_not_machine_names(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path, games=_games()))
    headers = [c.value for c in wb["Games"][4]]
    assert "Home Cover %" in headers
    assert "cover_probability" not in headers
    # A column with no entry in the table still reads sensibly.
    assert header_for("some_new_column") == "Some New Column"


def test_a_blank_export_cell_stays_blank(tmp_path: Path) -> None:
    """Ownership has no source, and a zero there would be a claim we never made."""
    wb = load_workbook(_build(tmp_path, dfs=_dfs()))
    ws = wb["DFS Pool"]
    headers = [c.value for c in ws[4]]
    own = headers.index("Ownership") + 1
    leverage = headers.index("Leverage") + 1
    assert ws.cell(row=5, column=own).value is None
    assert ws.cell(row=5, column=leverage).value is None


def test_percentages_are_formatted_as_percentages(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path, games=_games()))
    ws = wb["Games"]
    headers = [c.value for c in ws[4]]
    cover = ws.cell(row=5, column=headers.index("Home Cover %") + 1)
    assert cover.number_format == "0.0%"
    # ...and the underlying value is still the fraction, not a pre-multiplied 42.1
    assert cover.value == pytest.approx(0.421)


def test_the_curated_tiers_are_colour_coded(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path, plays=_plays()))
    ws = wb["Curated Plays"]
    a_plus = ws.cell(row=5, column=1)
    watch = ws.cell(row=6, column=1)
    assert a_plus.value == "A+"
    assert watch.value == "Watch"
    assert a_plus.fill.fgColor.rgb != watch.fill.fgColor.rgb


def test_the_dashboard_is_blocks_not_the_long_frame(tmp_path: Path) -> None:
    """There is no pivot on a tablet, so the sections are laid out here."""
    wb = load_workbook(_build(tmp_path, dashboard=_dashboard()))
    ws = wb["Dashboard"]
    text = [str(row[0].value or "") for row in ws.iter_rows(min_col=1, max_col=1)]
    for heading in ("How the model is doing", "Return on investment",
                    "Top plays", "Top props"):
        assert heading in text, f"missing block: {heading}"


def test_the_dashboard_survives_an_empty_record(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path))
    ws = wb["Dashboard"]
    text = [str(row[0].value or "") for row in ws.iter_rows(min_col=1, max_col=1)]
    assert "How the model is doing" in text
    assert any("nothing to report" in str(cell.value or "")
               for row in ws.iter_rows(min_col=1, max_col=3) for cell in row)


def test_read_me_states_the_sign_conventions_and_the_blanks(tmp_path: Path) -> None:
    """The two things a reader gets wrong silently if nobody tells them."""
    wb = load_workbook(_build(tmp_path, games=_games()))
    blob = " ".join(
        str(cell.value or "")
        for row in wb["Read Me"].iter_rows(min_col=1, max_col=2) for cell in row
    )
    assert "value on the HOME side" in blob
    assert "value on the OVER" in blob
    assert "Unknown, never zero" in blob
    assert "Seen and NOT bet" in blob


def test_read_me_counts_the_rows_in_the_file(tmp_path: Path) -> None:
    wb = load_workbook(_build(tmp_path, games=_games(), plays=_plays()))
    rows = {
        str(row[0].value): row[1].value
        for row in wb["Read Me"].iter_rows(min_col=1, max_col=2)
    }
    assert rows.get("Games") == 1
    assert rows.get("Curated Plays") == 2


def test_the_workbook_opens_after_a_round_trip_through_csv(tmp_path: Path) -> None:
    """The pipeline reads the CSVs back rather than rebuilding the frames."""
    csv_dir = tmp_path / "csv"
    csv_dir.mkdir()
    _games().to_csv(csv_dir / "games.csv", index=False, encoding="utf-8-sig")
    path = build_workbook(tmp_path / WORKBOOK_NAME, META,
                          games=pd.read_csv(csv_dir / "games.csv"))
    ws = load_workbook(path)["Games"]
    headers = [c.value for c in ws[4]]
    assert ws.cell(row=5, column=headers.index("Away") + 1).value == "Atlanta"
    assert ws.cell(row=5, column=headers.index("Proj Total") + 1) \
        .value == pytest.approx(44.2)


# ---------------------------------------------------------------------------
# Run status — the first thing an operator asks before kickoff.
# ---------------------------------------------------------------------------

def _readiness(verdict: str = "DEGRADED"):  # type: ignore[no-untyped-def]
    from velocity.export.readiness import assess

    one = pd.DataFrame([{"x": 1}])
    frames = {"games": pd.DataFrame([{"game_id": "g1",
                                      "kickoff": "2026-09-20T23:30:00Z"}]),
              "projections": one, "market": one, "plays": one,
              "props": None if verdict != "READY" else one,
              "team_totals": one,
              "dfs_pool": None if verdict != "READY" else one,
              "dfs_lineups": None if verdict != "READY" else one,
              "weather": one, "record": one}
    if verdict == "NOT READY":
        frames["market"] = None
    return assess(frames, now=pd.Timestamp("2026-09-20T22:51:00Z"),
                  generated_at=pd.Timestamp("2026-09-20T22:51:00Z"))


def test_run_status_leads_the_dashboard(tmp_path: Path) -> None:
    """Before "how is the model doing", answer "can I use this at all"."""
    path = build_workbook(tmp_path / WORKBOOK_NAME, META,
                          dashboard=_dashboard(), readiness=_readiness())
    ws = load_workbook(path)["Dashboard"]
    col_a = [str(r[0].value or "") for r in ws.iter_rows(min_col=1, max_col=1)]
    assert col_a.index("Run status") < col_a.index("How the model is doing")


def test_run_status_names_each_missing_surface(tmp_path: Path) -> None:
    path = build_workbook(tmp_path / WORKBOOK_NAME, META,
                          dashboard=_dashboard(), readiness=_readiness())
    ws = load_workbook(path)["Dashboard"]
    text = {str(r[0].value or ""): str(r[1].value or "")
            for r in ws.iter_rows(min_col=1, max_col=2)}
    assert text.get("Player props") == "missing"
    assert text.get("DFS pool") == "missing"
    assert text.get("Board (games)") == "ok"


def test_the_verdict_is_colour_coded(tmp_path: Path) -> None:
    fills = {}
    for verdict in ("READY", "DEGRADED", "NOT READY"):
        path = build_workbook(tmp_path / f"{verdict}.xlsx", META,
                              dashboard=_dashboard(), readiness=_readiness(verdict))
        ws = load_workbook(path)["Dashboard"]
        cell = next(r[0] for r in ws.iter_rows(min_col=1, max_col=1)
                    if str(r[0].value or "") == verdict)
        fills[verdict] = cell.fill.fgColor.rgb
    assert len(set(fills.values())) == 3, f"verdicts share a colour: {fills}"


def test_read_me_carries_the_verdict_too(tmp_path: Path) -> None:
    """Read Me is the tab that opens first."""
    path = build_workbook(tmp_path / WORKBOOK_NAME, META, readiness=_readiness())
    ws = load_workbook(path)["Read Me"]
    blob = " ".join(str(c.value or "") for r in ws.iter_rows(min_col=1, max_col=2)
                    for c in r)
    assert "Run status" in blob
    assert "DEGRADED" in blob


def test_a_workbook_without_readiness_still_builds(tmp_path: Path) -> None:
    """Readiness is additive: an older caller passing nothing must still work."""
    path = build_workbook(tmp_path / WORKBOOK_NAME, META, dashboard=_dashboard())
    ws = load_workbook(path)["Dashboard"]
    col_a = [str(r[0].value or "") for r in ws.iter_rows(min_col=1, max_col=1)]
    assert "Run status" not in col_a
    assert "How the model is doing" in col_a


# ---------------------------------------------------------------------------
# Bet Card — the one tab built for a phone held upright (Phase 13 Stage 1).
# ---------------------------------------------------------------------------

def _card_plays() -> pd.DataFrame:
    return pd.DataFrame([
        {"tier": "B", "bet_type": "prop", "selection": "Bijan Over 64.5 rush_yds",
         "market": "rush_yds", "edge": 0.031, "confidence": 3.7, "stake": 0.93,
         "reason": "Bijan Over 64.5 rush_yds — model 60.0% · fair 53.0%"},
        {"tier": "Watch", "bet_type": "game", "selection": "Miami TT Over 16.5",
         "market": "team_total_away", "edge": 0.034, "confidence": 4.3,
         "stake": 0.0, "reason": "Miami TT Over 16.5 — VETOED by the intel layer"},
        {"tier": "A+", "bet_type": "game", "selection": "Under 41.5",
         "market": "total", "edge": 0.053, "confidence": 5.8, "stake": 1.69,
         "reason": "Under 41.5 — Model total 37 · rule A (unders 4+)"},
    ])


def _card_text(path: Path) -> list[str]:
    ws = load_workbook(path)["Bet Card"]
    return [str(c.value) for row in ws.iter_rows(min_col=1, max_col=3)
            for c in row if c.value not in (None, "")]


def test_the_bet_card_leads_with_the_plays_that_get_bet(tmp_path: Path) -> None:
    """A+ before A before B, and Watch last under its own heading."""
    text = _card_text(build_workbook(tmp_path / WORKBOOK_NAME, META,
                                     plays=_card_plays()))
    assert text.index("A+") < text.index("B") < text.index("Watch")
    assert any("Watch — seen, not bet" in t for t in text)


def test_the_bet_card_keeps_the_declined_plays(tmp_path: Path) -> None:
    """A card that silently drops them reads like a system that never looked."""
    text = _card_text(build_workbook(tmp_path / WORKBOOK_NAME, META,
                                     plays=_card_plays()))
    assert any("Miami TT Over 16.5" in t for t in text)


def test_the_bet_card_does_not_repeat_the_selection_in_the_reason(tmp_path: Path) -> None:
    """Line 1 is the call; line 3 is the argument, not the call again."""
    ws = load_workbook(build_workbook(tmp_path / WORKBOOK_NAME, META,
                                      plays=_card_plays()))["Bet Card"]
    reasons = [str(c.value) for row in ws.iter_rows(min_col=2, max_col=2)
               for c in row if c.value and "rule A" in str(c.value)]
    assert reasons and not reasons[0].startswith("Under 41.5")


def test_the_bet_card_wraps_its_argument(tmp_path: Path) -> None:
    """Portrait: the reason wraps rather than running off the screen."""
    ws = load_workbook(build_workbook(tmp_path / WORKBOOK_NAME, META,
                                      plays=_card_plays()))["Bet Card"]
    wrapped = [c for row in ws.iter_rows(min_col=2, max_col=2) for c in row
               if c.value and c.alignment and c.alignment.wrap_text]
    assert wrapped, "no wrapped reason cell"


def test_the_bet_card_says_so_when_there_is_nothing(tmp_path: Path) -> None:
    text = _card_text(build_workbook(tmp_path / WORKBOOK_NAME, META))
    assert any("No plays" in t for t in text)


def test_the_two_card_tabs_are_not_confusable(tmp_path: Path) -> None:
    """'Bet Card' and 'Betting Card' side by side was a trap; the wide
    one-row-per-game table is now named for what it is."""
    names = load_workbook(build_workbook(tmp_path / WORKBOOK_NAME, META)).sheetnames
    assert "Games" in names
    assert "Betting Card" not in names


def test_the_bet_card_leads_with_the_game_not_the_market(tmp_path) -> None:
    """The fix for "we don't know opponents".

    Line 1 is the call and the stake; line 2 used to open with the market,
    which the selection had usually already said. It now opens with the
    matchup and the kickoff — the two things a prop row cannot tell you and
    which decide whether the card is still worth acting on.
    """
    plays = pd.DataFrame([{
        "tier": "A", "bet_type": "prop",
        "selection": "Matthew Golden Under 4.5 receptions",
        "matchup": "Green Bay Packers @ Atlanta Falcons",
        "kickoff": "2026-09-22T00:15:00Z",
        "market": "receptions", "edge": 0.0895, "confidence": 9.4,
        "stake": 1.35, "reason": "model 64.5% · fair 55.5%",
    }])
    dest = build_workbook(
        tmp_path / "velocity.xlsx",
        ExportMeta(generated_at="2026-09-21T16:38:49Z", season=2026, week=3),
        plays=plays,
    )

    sheet = load_workbook(dest)["Bet Card"]
    line2 = [c for row in sheet.iter_rows(min_row=4, max_row=6, values_only=True)
             for c in row if c and "receptions" in str(c) and "edge" in str(c)]
    assert line2, "the card should carry a numbers line"
    assert line2[0].startswith("Green Bay Packers @ Atlanta Falcons · Mon 7:15P CT · ")


def test_a_kickoff_the_card_cannot_read_is_simply_absent(tmp_path) -> None:
    """No crash, no epoch, no "Invalid Date" — the label is just dropped."""
    plays = pd.DataFrame([{
        "tier": "B", "bet_type": "game", "selection": "Under 59.5",
        "matchup": "Green Bay Packers @ Atlanta Falcons", "kickoff": "not a date",
        "market": "total", "edge": 0.044, "confidence": 5.5,
        "stake": 1.32, "reason": "Model total 51",
    }])
    dest = build_workbook(
        tmp_path / "velocity.xlsx",
        ExportMeta(generated_at="2026-09-21T16:38:49Z", season=2026, week=3),
        plays=plays,
    )

    sheet = load_workbook(dest)["Bet Card"]
    line2 = [c for row in sheet.iter_rows(min_row=4, max_row=6, values_only=True)
             for c in row if c and "edge" in str(c)]
    assert line2[0] == "Green Bay Packers @ Atlanta Falcons · total · edge 4.4% · confidence 5.5"
